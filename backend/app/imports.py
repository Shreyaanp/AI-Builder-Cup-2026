"""Import a menu or sales export from any POS or spreadsheet (CSV / Excel).

No per-POS integration: owners upload whatever their POS exports. Gemini maps the columns
(item, price, unit, quantity, date, stock); offline mode uses header heuristics. A menu becomes
a sealed inventory proposal the owner confirms; sales are sealed separately after a summary.
"""

from __future__ import annotations

import csv
import io
import itertools
import re
from collections import defaultdict
from datetime import date, datetime
from typing import Any, Literal

import structlog
from pydantic import BaseModel, Field

from .settings import Settings

log = structlog.get_logger()

MAX_BYTES = 2_000_000
MAX_ROWS = 2_000
SAMPLE_ROWS = 15
CSV_TYPES = {"text/csv", "application/csv", "text/plain"}
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ImportRejected(Exception):
    def __init__(self, status: int, code: str) -> None:
        super().__init__(code)
        self.status = status
        self.code = code


class ColumnMap(BaseModel):
    kind: Literal["menu", "sales"] = Field(
        description="menu: a list of items and prices; sales: rows of sold items over time"
    )
    name: int | None = Field(default=None, description="0-based column index of the item name")
    price: int | None = Field(default=None, description="Unit price column")
    unit: int | None = Field(default=None, description="Selling unit column, if any")
    quantity: int | None = Field(default=None, description="Quantity sold column (sales)")
    total: int | None = Field(default=None, description="Line total / revenue column (sales)")
    date: int | None = Field(default=None, description="Date or time of sale column (sales)")
    stock: int | None = Field(default=None, description="Stock on hand column (menu)")
    category: int | None = Field(default=None, description="Item category or menu section column")


PROMPT = """You map a spreadsheet exported by a shop's POS or written by the owner. Given the
header row and sample rows, decide whether it is a menu/price list (kind "menu") or a sales or
transactions export (kind "sales"), and give the 0-based column index for each role. Use null
for roles that are not present. Never guess a column whose contents don't fit the role."""

HEADER_HINTS: dict[str, tuple[str, ...]] = {
    "category": ("category", "group", "section", "department"),
    "name": ("item", "product", "name", "description", "menu", "dish", "sku name", "article"),
    "price": ("price", "unit price", "selling price", "rate", "cost", "amount"),
    "unit": ("unit", "uom", "size", "per"),
    "quantity": ("qty", "quantity", "sold", "count", "units sold"),
    "total": ("total", "net sales", "gross sales", "revenue", "line total", "subtotal"),
    "date": ("date", "time", "created", "receipt date", "day"),
    "stock": ("stock", "inventory", "on hand", "in stock", "available"),
}


def parse_table(data: bytes, content_type: str) -> tuple[list[str], list[list[str]]]:
    """Return the header row and up to MAX_ROWS data rows as strings."""
    if len(data) > MAX_BYTES:
        raise ImportRejected(413, "FILE_TOO_LARGE")
    if content_type == XLSX_TYPE or data[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        except Exception:
            raise ImportRejected(415, "UNREADABLE_FILE") from None
        if sheet is None:
            raise ImportRejected(422, "EMPTY_FILE")
        rows = [
            ["" if cell is None else _text(cell) for cell in row]
            for row in sheet.iter_rows(values_only=True, max_row=MAX_ROWS + 2)
        ]
    elif content_type in CSV_TYPES or not content_type:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        rows = list(itertools.islice(reader, MAX_ROWS + 2))
    else:
        raise ImportRejected(415, "UNSUPPORTED_FILE")
    if len(rows) > MAX_ROWS + 1:
        raise ImportRejected(413, "TOO_MANY_ROWS")
    rows = [row for row in rows if any(cell.strip() for cell in row)]
    if len(rows) < 2:
        raise ImportRejected(422, "EMPTY_FILE")
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    return [cell.strip() for cell in rows[0]], [[c.strip() for c in row] for row in rows[1:]]


def _text(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def heuristic_map(headers: list[str]) -> ColumnMap:
    found: dict[str, int | None] = {}
    lowered = [header.casefold().strip() for header in headers]
    used: set[int] = set()
    # Exact header matches win, a column serves one role, and "unit" goes last so that a
    # "Unit Price" column is taken as the price rather than the unit.
    for role in ("name", "price", "quantity", "total", "date", "stock", "category", "unit"):
        hints = HEADER_HINTS[role]
        free = [(index, header) for index, header in enumerate(lowered) if index not in used]
        match = next((index for index, header in free if header in hints), None)
        if match is None:
            match = next((index for index, header in free if any(h in header for h in hints)), None)
        found[role] = match
        if match is not None:
            used.add(match)
    # "Amount" alone is a price on a menu but a total on a sales export.
    is_sales = found["date"] is not None or found["quantity"] is not None
    if is_sales and found["total"] is None and found["price"] is not None:
        if "amount" in lowered[found["price"]]:
            found["total"], found["price"] = found["price"], None
    return ColumnMap(kind="sales" if is_sales else "menu", **found)


def number(text: str) -> float | None:
    cleaned = re.sub(r"[^\d.,\-]", "", text or "")
    if not cleaned:
        return None
    if "," in cleaned and "." in cleaned:
        cleaned = cleaned.replace(",", "")
    elif cleaned.count(",") == 1 and len(cleaned.split(",")[1]) in (1, 2):
        cleaned = cleaned.replace(",", ".")
    else:
        cleaned = cleaned.replace(",", "")
    try:
        return float(cleaned)
    except ValueError:
        return None


def cell(row: list[str], index: int | None) -> str:
    return row[index] if index is not None and 0 <= index < len(row) else ""


def menu_items(rows: list[list[str]], mapping: ColumnMap) -> tuple[list[dict[str, Any]], int]:
    """Unique (name, unit) items with a major-unit price; returns items and skipped rows."""
    items: dict[tuple[str, str], dict[str, Any]] = {}
    skipped = 0
    for row in rows:
        name = cell(row, mapping.name)[:180]
        price = number(cell(row, mapping.price))
        if not name or price is None or price < 0:
            skipped += 1
            continue
        unit = (cell(row, mapping.unit) or "item")[:40]
        stock = number(cell(row, mapping.stock))
        items[(name.casefold(), unit.casefold())] = {
            "name": name,
            "price": price,
            "unit": unit,
            "stock": int(stock) if stock is not None and stock >= 0 else None,
            **(
                {"category": cell(row, mapping.category)[:80]}
                if cell(row, mapping.category)
                else {}
            ),
        }
    return list(items.values())[:200], skipped


def sales_summary(rows: list[list[str]], mapping: ColumnMap) -> dict[str, Any]:
    quantity: dict[str, float] = defaultdict(float)
    revenue: dict[str, float] = defaultdict(float)
    days: set[str] = set()
    for row in rows:
        name = cell(row, mapping.name)
        if not name:
            continue
        qty = number(cell(row, mapping.quantity)) if mapping.quantity is not None else 1
        qty = qty if qty is not None else 0
        total = number(cell(row, mapping.total))
        if total is None:
            total = (number(cell(row, mapping.price)) or 0) * qty
        quantity[name] += qty
        revenue[name] += total
        if day := cell(row, mapping.date)[:10]:
            days.add(day)
    top = sorted(quantity, key=lambda item: (-quantity[item], item))[:10]
    return {
        "rows": len(rows),
        "items": len(quantity),
        "days": len(days),
        "first_day": min(days) if days else None,
        "last_day": max(days) if days else None,
        "total_revenue": round(sum(revenue.values()), 2),
        "top_items": [
            {"name": name, "quantity": quantity[name], "revenue": round(revenue[name], 2)}
            for name in top
        ],
    }


class ColumnMapper:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings
        self._client: Any = None

    def client(self) -> Any:
        from google import genai

        if self._client is None:
            self._client = (
                genai.Client(
                    vertexai=True,
                    project=self.settings.google_cloud_project,
                    location=self.settings.google_cloud_location,
                )
                if self.settings.gemini_backend == "vertex"
                else genai.Client(api_key=self.settings.google_api_key.get_secret_value())
            )
        return self._client

    async def map(self, headers: list[str], sample: list[list[str]]) -> ColumnMap:
        fallback = heuristic_map(headers)
        if self.settings.coach_mode != "gemini":
            return fallback
        try:
            from google.genai import types

            table = "\n".join(" | ".join(row) for row in [headers, *sample[:SAMPLE_ROWS]])
            response = await self.client().aio.models.generate_content(
                model=self.settings.google_text_model,
                contents=table[:8000],
                config=types.GenerateContentConfig(
                    system_instruction=PROMPT,
                    response_mime_type="application/json",
                    response_schema=ColumnMap,
                    temperature=0,
                    max_output_tokens=300,
                ),
            )
            mapped = response.parsed
            if isinstance(mapped, ColumnMap) and mapped.name is not None:
                return mapped
        except Exception:
            log.warning("column_map_failed")
        return fallback

    async def close(self) -> None:
        if self._client is not None:
            await self._client.aio.aclose()
